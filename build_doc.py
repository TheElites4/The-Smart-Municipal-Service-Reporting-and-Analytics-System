# build_doc.py - Builds MzansiConnect-Documentation.docx using python-docx.
# Embeds the captured Android emulator screenshots and reserves captioned
# placeholders for the web dashboards and PostgreSQL screenshots the author
# will supply. Run from the repository root:  python build_doc.py

import os
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

ROOT = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(ROOT, "MzansiConnect-Documentation.docx")

doc = Document()

# Base styles
normal = doc.styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(10.5)
_normal_pf = normal.paragraph_format
_normal_pf.space_before = Pt(0)
_normal_pf.space_after = Pt(4)
_normal_pf.line_spacing = 1.0

for _name in ("Heading 1", "Heading 2", "Heading 3"):
    _s = doc.styles[_name]
    _s.paragraph_format.space_before = Pt(8 if _name == "Heading 1" else 6)
    _s.paragraph_format.space_after = Pt(3)
    _s.paragraph_format.line_spacing = 1.0

for section in doc.sections:
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)


def heading(text, level):
    h = doc.add_heading(text, level=level)
    return h


def body(text):
    p = doc.add_paragraph(text)
    p.paragraph_format.space_after = Pt(6)
    return p


def centered(text, bold=False, italic=False, size=11):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    r.bold = bold
    r.italic = italic
    r.font.size = Pt(size)
    return p


def caption(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    r.italic = True
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor(0x40, 0x40, 0x40)
    return p


def add_image(filename, cap, width_in=2.6):
    path = os.path.join(ROOT, filename)
    if os.path.exists(path):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(path, width=Inches(width_in))
        caption(cap)
    else:
        caption(f"[ INSERT SCREENSHOT HERE - {cap} (missing {filename}) ]")


def image_grid(items, width_in=0.95, cols=4):
    """Place figures in a multi-column borderless table to save vertical space."""
    rows = (len(items) + cols - 1) // cols
    t = doc.add_table(rows=rows, cols=cols)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for idx, (filename, cap) in enumerate(items):
        cell = t.cell(idx // cols, idx % cols)
        cell.text = ""
        path = os.path.join(ROOT, filename)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if os.path.exists(path):
            p.add_run().add_picture(path, width=Inches(width_in))
        cp = cell.add_paragraph()
        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = cp.add_run(cap)
        r.italic = True
        r.font.size = Pt(8)
    doc.add_paragraph()
    return t


def placeholder(cap):
    caption(f"[ INSERT SCREENSHOT HERE - {cap} ]")


def table(rows):
    t = doc.add_table(rows=len(rows), cols=len(rows[0]))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            cell = t.cell(r, c)
            cell.text = ""
            run = cell.paragraphs[0].add_run(val)
            run.font.size = Pt(9)
            if r == 0:
                run.bold = True
    doc.add_paragraph()
    return t


# ---------------------------------------------------------------------------
# Title
# ---------------------------------------------------------------------------
centered("MzansiConnect", bold=True, size=20)
centered("Smart Municipal Service Reporting and Analytics System", size=14)
centered("Technical and Academic Documentation", italic=True, size=11)
centered("Version 1.0  |  2026  |  Windows 11 (x64), Node.js, PostgreSQL 16, Android", size=10)

# ---------------------------------------------------------------------------
# Abstract
# ---------------------------------------------------------------------------
heading("Abstract", 1)
body(
    "MzansiConnect is an integrated municipal service-reporting platform that connects citizens "
    "to local-government service delivery. The system comprises four presentation applications - "
    "a React citizen web portal, an Android citizen client, a staff operations dashboard, and an "
    "administrative analytics dashboard - that all read from and write to a single "
    "PostgreSQL-backed Node.js API. This document describes the purpose, architecture, and "
    "interface of each application, presents captured screenshots of the working system, and "
    "documents the database layer that acts as the authoritative source of truth. The report "
    "serves both as a technical reference and as an academic record of the implemented system."
)
p = body(
    "Keywords: service delivery, e-government, PostgreSQL, REST API, Android, React, database "
    "normalisation, multi-factor authentication."
)
p.runs[0].italic = True

# ---------------------------------------------------------------------------
# 1. Introduction
# ---------------------------------------------------------------------------
heading("1. Introduction", 1)
heading("1.1 Background", 2)
body(
    "South African municipalities frequently receive service-delivery complaints through "
    "unstructured channels such as telephone calls, paper forms, and social media, which makes "
    "prioritisation, tracking, and reporting difficult. MzansiConnect addresses this by providing "
    "one structured pipeline: a citizen submits a report with a category, municipality, "
    "description, and location; the report is committed to PostgreSQL; and it then appears "
    "immediately in municipality-scoped staff queues and administrative analytics."
)
heading("1.2 Objectives", 2)
for t in [
    "1. Provide a single, validated intake channel for municipal service reports.",
    "2. Store every report as an authoritative database row with an SLA deadline.",
    "3. Give staff a municipality-scoped queue with controlled status transitions.",
    "4. Provide administrators with database-backed analytics and workforce management.",
    "5. Secure staff and administrative access using password plus one-time MFA codes.",
]:
    body(t)
heading("1.3 Scope", 2)
body(
    "The system delivers four applications and one shared API against one PostgreSQL database. "
    "Sample analytics and sample outage rows were deliberately removed so that all dashboards "
    "reflect real committed data only."
)

# ---------------------------------------------------------------------------
# 2. Architecture
# ---------------------------------------------------------------------------
heading("2. System Architecture", 1)
heading("2.1 Component Overview", 2)
table([
    ["Component", "Technology", "Role"],
    ["Citizen Portal", "Vite + React (TypeScript)", "Citizen web intake and tracking"],
    ["Android Client", "Kotlin, Jetpack Compose, Room", "Native citizen mobile intake and tracking"],
    ["Staff Dashboard", "Static HTML/CSS/JS", "Municipality-scoped operations queue"],
    ["Admin Dashboard", "Static HTML/CSS/JS", "Workforce management and analytics"],
    ["API", "Node.js, pg", "Validation, persistence, auth, realtime events"],
    ["Database", "PostgreSQL 16", "Single source of truth"],
])
heading("2.2 Data Flow", 2)
body(
    "A citizen submission calls POST /api/reports. The API validates the category and municipality "
    "against PostgreSQL reference tables, commits the row (including the 48-hour default "
    "sla_due_at), then broadcasts a server-sent event on GET /api/events. Staff and administrative "
    "clients react by reloading GET /api/dashboard; the Android client refreshes its Room mirror "
    "only after the API confirms. Room is a replaceable display cache, never the system of record."
)
heading("2.3 Authentication", 2)
body(
    "Dashboard login is a two-step process. POST /api/auth/login verifies the password and "
    "dispatches a one-time code through Resend (email) or Twilio (SMS). POST /api/auth/verify-mfa "
    "validates the hashed, ten-minute, five-attempt code and issues a Bearer token. Provider "
    "credentials remain server-side; the web and Android clients never hold provider secrets."
)

# ---------------------------------------------------------------------------
# 3. Citizen Portal
# ---------------------------------------------------------------------------
heading("3. Citizen Portal (Web)", 1)
body(
    "The citizen web portal is a Vite/React single-page application served on "
    "http://localhost:5173. It loads its category and municipality choices from GET /api/options, "
    "submits reports, and renders each citizen's own reports by stable external reference. "
    "Interface features include a hero landing section, a guided report form with GPS "
    "reverse-geocoding, a load-shedding banner, a community feed, a citizen assistant, and a "
    "gamified leaderboard."
)
placeholder("Figure 3.1: Citizen portal landing and service-report entry")

# ---------------------------------------------------------------------------
# 4. Android
# ---------------------------------------------------------------------------
heading("4. Android Citizen Application", 1)
body(
    "The Android client (MzansiConnect/) is a native Kotlin/Compose application that posts to "
    "POST /api/reports and mirrors confirmed rows into a local Room database. It identifies the "
    "citizen through a stable app_users.external_ref. The following figures were captured from the "
    "running application in an Android emulator using http://10.0.2.2:8787/ as the API base URL."
)
image_grid([
    ("emu-step1.png", "Figure 4.1 - Application launch and splash"),
    ("emu-screen.png", "Figure 4.2 - Citizen home and navigation"),
    ("emu-screen3.png", "Figure 4.3 - Service report form"),
    ("emu-screen4.png", "Figure 4.4 - Category and municipality selection"),
    ("emu-screen6.png", "Figure 4.5 - Report submission confirmation"),
    ("emu-screen8.png", "Figure 4.6 - Report list and status tracking"),
    ("emu-screen9.png", "Figure 4.7 - Report detail with SLA countdown"),
])

# ---------------------------------------------------------------------------
# 5. Staff Dashboard
# ---------------------------------------------------------------------------
heading("5. Staff Operations Dashboard", 1)
body(
    "The staff dashboard is a static client served at http://localhost:8080/Staff%20Dashboard/. "
    "After MFA login it presents a municipality-scoped Service queue (Open work). Each job card "
    "shows the public reference, category, urgency, description, location, citizen chip, and "
    "current status, with a status selector supporting the transitions Submitted -> Under review "
    "-> Assigned -> In progress -> Resolved -> Closed. A search box filters the queue by "
    "reference, category, or location, and a live-status indicator reflects the server-sent event "
    "connection. A load-shedding panel queries the backend's EskomSePush integration."
)
placeholder("Figure 5.1: Staff sign-in screen")
placeholder("Figure 5.2: Staff service queue with status controls")

# ---------------------------------------------------------------------------
# 6. Admin Dashboard
# ---------------------------------------------------------------------------
heading("6. Admin Dashboard", 1)
body(
    "The admin dashboard is served at http://localhost:8080/Admin%20Dashboard/ and is protected "
    "by the same password-plus-MFA flow with the admin role. It presents a Municipal pulse header "
    "with API-monitoring health, an overview statistic grid, a load-shedding panel, a Workforce "
    "panel (active employees and open assignments), and a Recent demand panel (highest-priority "
    "reports first). It also provides the Add staff member form that calls POST /api/staff, "
    "hashing the password with PostgreSQL crypt() and recording municipality, department, and an "
    "optional SMS number. Administrative analytics are drawn from GET /api/admin/analytics, which "
    "aggregates category, status, monthly, and municipality counts from service_reports."
)
placeholder("Figure 6.1: Admin sign-in screen")
placeholder("Figure 6.2: Admin municipal pulse, workforce, and demand panels")

# ---------------------------------------------------------------------------
# 7. Database
# ---------------------------------------------------------------------------
heading("7. Database Layer (PostgreSQL)", 1)
heading("7.1 Schema", 2)
body(
    "The schema (backend/schema.sql) defines enumerated types user_role, report_status, and "
    "urgency_level, and the tables app_users, service_categories, service_reports, "
    "report_assignments, report_status_events, audit_events, and auth_mfa_challenges. Queue, "
    "citizen, assignment, workforce, and MFA indexes support the dashboard access patterns, and a "
    "touch_updated_at() trigger maintains updated_at on update."
)
heading("7.2 Municipal (Data Normalisation) Migration", 2)
body(
    "Staff accounts are created through an administrative free-text Municipality field, so the "
    "same municipality may be stored in several textual forms (sol plaatje municipality, Sol "
    "Plaatje Local Municipality, City of Johannesburg, Johannesburg). Without normalisation, a "
    "staff member can sign in and see an empty queue. Migration 003_normalise_municipality.sql "
    "introduces the normalise_municipality(value) function, declared IMMUTABLE, which lower-cases, "
    "trims, collapses whitespace, strips an optional City of/Municipality of prefix, and strips an "
    "optional trailing Local/Metropolitan/District Municipality or Municipality suffix. It mirrors "
    "normaliseMunicipality() in the API so that application and database comparisons agree, and it "
    "is paired with the functional index reports_municipality_norm_idx so municipality-scoped "
    "queries remain fast."
)
placeholder("Figure 7.1: PostgreSQL normalise_municipality applied to stored reports")
heading("7.3 Reference Data", 2)
body(
    "Service categories are seeded in service_categories, and the municipality directory is "
    "populated by the explicit importer (npm run import:municipalities), which aborts unless it "
    "reads exactly 8 metropolitan, 44 district, and 205 local municipalities - it never publishes "
    "a partial list. Optional report photos (migration 006) are stored in report_photos and served "
    "only to the submitting citizen or to municipality-scoped staff/admin accounts."
)
placeholder("Figure 7.2: Municipality reference table in PostgreSQL")
placeholder("Figure 7.3: Service reports with normalised municipality values")

# ---------------------------------------------------------------------------
# 8. Evaluation
# ---------------------------------------------------------------------------
heading("8. Evaluation and Discussion", 1)
body(
    "The architecture centralises authority in PostgreSQL, which removes the divergence that "
    "arises when multiple clients hold their own state. Key strengths are: validated, structured "
    "intake; a single SLA-authoritative deadline; role-scoped tokens enforced by the API; and "
    "analytics computed only from real rows. The principal operational dependency is the MFA "
    "provider: seeded example.gov.za addresses are placeholders and cannot receive codes, so "
    "production access requires real Resend/Twilio credentials and reachable contacts. For "
    "production the system requires HTTPS, a managed secret store, database backups, rate "
    "limiting, and an established identity provider."
)

# ---------------------------------------------------------------------------
# 9. Conclusion
# ---------------------------------------------------------------------------
heading("9. Conclusion", 1)
body(
    "MzansiConnect demonstrates a coherent multi-client service-delivery platform in which a React "
    "web portal, a native Android client, a staff operations dashboard, and an administrative "
    "analytics dashboard share one PostgreSQL system of record. The documented migrations - in "
    "particular the municipality normalisation function - illustrate how application-level and "
    "database-level agreement on data semantics is essential to correct, non-empty staff queues and "
    "reliable analytics."
)

# ---------------------------------------------------------------------------
# Appendices
# ---------------------------------------------------------------------------
heading("Appendix A - API Endpoints", 1)
table([
    ["Method", "Path", "Access", "Purpose"],
    ["GET", "/health", "Public", "API and PostgreSQL health"],
    ["GET", "/api/options", "Public", "Categories and municipalities"],
    ["POST", "/api/auth/login", "Public", "Check password and send MFA code"],
    ["POST", "/api/auth/verify-mfa", "Public", "Verify code and issue token"],
    ["GET", "/api/dashboard", "Staff/admin", "Report queue and workforce summary"],
    ["GET", "/api/admin/analytics", "Admin", "Report aggregates from PostgreSQL"],
    ["POST", "/api/reports", "Citizen", "Commit a report and broadcast update"],
    ["GET", "/api/reports?externalRef=..", "Citizen", "Read a citizen's reports"],
    ["PATCH", "/api/reports/:id/status", "Staff/admin", "Update status and audit trail"],
    ["GET", "/api/events", "All roles", "Server-sent live updates"],
])
heading("Appendix B - Migration Order", 1)
body(
    "1. schema.sql   2. 002_external_ref.sql   3. 003_normalise_municipality.sql   "
    "4. 004_mfa_challenges.sql   5. 005_service_categories.sql   "
    "6. 006_municipalities_and_photos.sql   7. npm run import:municipalities   8. seed.sql"
)

doc.save(OUT)
print(f"Created: {OUT}")
